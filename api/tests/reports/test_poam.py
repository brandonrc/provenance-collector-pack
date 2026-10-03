import csv
import io
from datetime import timedelta

from openpyxl import load_workbook

from conftest import NOW, VULNS, make_snapshot
from posture.reports.poam import EMASS_COLUMNS, GENERIC_COLUMNS, LEGACY_COLUMNS
from posture.reports.registry import generate

N_FINDINGS = sum(len(v[-1]) for v in VULNS)


def _failing_checks(snapshot):
    return {r.check_id for r in snapshot.posture_results if r.status == "fail"}


def _wb(snapshot, opts, **o):
    rep = generate("poam", "xlsx", snapshot, {**opts, **o})
    assert rep.content_type.startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(rep.content))


def test_xlsx_sheets_and_headers(snapshot, opts):
    wb = _wb(snapshot, opts)
    assert wb.sheetnames == ["eMASS", "POA&M", "eMASS (legacy)", "Info"]
    assert [c.value for c in wb["eMASS"][1]] == EMASS_COLUMNS
    assert [c.value for c in wb["POA&M"][1]] == GENERIC_COLUMNS
    assert [c.value for c in wb["eMASS (legacy)"][1]] == LEGACY_COLUMNS
    assert len(EMASS_COLUMNS) == 32 and len(GENERIC_COLUMNS) == 27


def test_row_per_image_cve_package(snapshot, opts):
    wb = _wb(snapshot, opts)
    expected = N_FINDINGS + len(_failing_checks(snapshot))
    for name in ("eMASS", "POA&M", "eMASS (legacy)"):
        assert wb[name].max_row - 1 == expected, name
    ids = [r[0].value for r in wb["POA&M"].iter_rows(min_row=2)]
    assert len(ids) == len(set(ids)), "POAM IDs must be unique"
    assert any(i.startswith("SP-CFG-") for i in ids)


def test_rollup_by_cve(snapshot, opts):
    wb = _wb(snapshot, opts, rollupByCve=True)
    n_cves = len({v[0] for v in VULNS})
    assert wb["eMASS"].max_row - 1 == n_cves + len(_failing_checks(snapshot))
    ws = wb["POA&M"]
    hdr = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr.index("POAM ID")] == "SP-CVE-2024-24790")
    assets = row[hdr.index("Asset Identifier")]
    assert "postgres" in assets and "coredns" in assets  # rolled up across two images


def test_sla_and_fields(snapshot, opts):
    ws = _wb(snapshot, opts)["POA&M"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    r = next(x for x in rows if x["Weakness Source Identifier"] == "CVE-2024-6387")
    detected = (NOW - timedelta(days=40)).date()
    assert r["Original Detection Date"].date() == detected
    assert r["Scheduled Completion Date"].date() == detected + timedelta(days=15)  # critical SLA
    assert "OVERDUE" in r["Comments"]
    assert r["Controls"] == "RA-5, SI-2, SI-2(2)"
    assert r["Original Risk Rating"] == "High"
    assert r["Vendor Dependency"] == "No"
    nofix = next(x for x in rows if x["Weakness Source Identifier"] == "CVE-2023-52425")
    assert nofix["Vendor Dependency"] == "Yes" and nofix["Vendor Dependent Product Name"] == "libexpat1"
    assert nofix["Controls"] == "RA-5, SI-2"
    priv = next(x for x in rows if x["POAM ID"] == "SP-CFG-privileged")
    assert priv["Controls"] == "AC-6, CM-7"
    assert "default/Deployment/legacy-proxy" in priv["Asset Identifier"]
    assert "V-233127" in priv["Weakness Source Identifier"]
    # rows are sorted most severe first
    assert rows[0]["Original Risk Rating"] == "High"


def test_emass_sheet_values(snapshot, opts):
    ws = _wb(snapshot, opts)["eMASS"]
    hdr = [c.value for c in ws[1]]
    rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    r = next(x for x in rows if x["Security Checks"] == "CVE-2024-45490")
    assert r["POA&M Status"] == "Ongoing"
    assert r["POA&M Item ID"] in (None, "")
    assert r["Controls / APs"] == "RA-5"
    assert r["Raw Severity"] == "Very High" and r["Severity"] == "Very High"
    assert r["Milestone ID"] == 1 and r["Milestone Status"] == "Pending"
    assert r["POA&M Scheduled Completion Date"] == r["Milestone Scheduled Completion Date"]
    assert "nginx" in r["Devices Affected"]
    assert ws.cell(row=2, column=hdr.index("POA&M Scheduled Completion Date") + 1).number_format == "mm/dd/yyyy"
    cfg = next(x for x in rows if (x["Security Checks"] or "").startswith("no-netpol"))
    assert cfg["Controls / APs"] == "SC-7"
    assert cfg["Identification Source"].startswith("Container Platform Security Requirements Guide")


def test_csv_variants(snapshot, opts):
    for variant, cols in (("emass", EMASS_COLUMNS), ("generic", GENERIC_COLUMNS), ("emass-legacy", LEGACY_COLUMNS)):
        rep = generate("poam", "csv", snapshot, {**opts, "poamVariant": variant})
        text = rep.content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        assert rows[0] == cols
        assert len(rows) - 1 == N_FINDINGS + len(_failing_checks(snapshot))
        assert all(len(r) == len(cols) for r in rows)
    rows = list(csv.DictReader(io.StringIO(generate("poam", "csv", snapshot, opts).content.decode("utf-8-sig"))))
    assert rows[0]["POA&M Scheduled Completion Date"].count("/") == 2  # MM/DD/YYYY


def test_scope_namespace(opts):
    snap = make_snapshot(scope={"kind": "namespace", "name": "dev"})
    ws = _wb(snap, opts)["POA&M"]
    text = " ".join(str(c.value) for row in ws.iter_rows(min_row=2) for c in row)
    assert "nginx" not in text and "jupyterhub" in text
    assert "legacy-proxy" not in text


def test_exclude_system_namespaces(snapshot, opts):
    ws = _wb(snapshot, opts, includeSystemNamespaces=False)["POA&M"]
    text = " ".join(str(c.value) for row in ws.iter_rows(min_row=2) for c in row)
    assert "coredns" not in text and "kube-system" not in text
