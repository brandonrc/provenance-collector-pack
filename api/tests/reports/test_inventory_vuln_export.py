import csv
import io
import json

from openpyxl import load_workbook

from conftest import IMAGES, VULNS
from posture.reports.inventory import IMAGE_COLUMNS
from posture.reports.registry import generate
from posture.reports.vuln_export import CSV_COLUMNS


def test_inventory_xlsx(snapshot, opts):
    wb = load_workbook(io.BytesIO(generate("inventory", "xlsx", snapshot, opts).content))
    assert wb.sheetnames == ["Images", "Workloads", "Namespaces", "Info"]
    ws = wb["Images"]
    assert [c.value for c in ws[1]] == IMAGE_COLUMNS
    assert ws.max_row - 1 == len(IMAGES)
    hdr = IMAGE_COLUMNS
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr.index("Image Reference")].startswith("docker.io/library/nginx"))
    assert row[hdr.index("Digest")].startswith("sha256:")
    assert row[hdr.index("Base OS")] == "debian 12.4"
    assert row[hdr.index("Scanner Coverage")] == "3/3"
    assert wb["Workloads"].max_row - 1 == 7


def test_inventory_csv(snapshot, opts):
    rows = list(csv.reader(io.StringIO(generate("inventory", "csv", snapshot, opts).content.decode("utf-8-sig"))))
    assert rows[0] == IMAGE_COLUMNS and len(rows) == len(IMAGES) + 1


def test_vuln_csv(snapshot, opts):
    rows = list(csv.DictReader(io.StringIO(generate("vuln-export", "csv", snapshot, opts).content.decode("utf-8-sig"))))
    assert list(rows[0].keys()) == CSV_COLUMNS
    assert len(rows) == sum(len(v[-1]) for v in VULNS)
    r = next(x for x in rows if x["Vulnerability ID"] == "CVE-2023-44487")
    assert r["Trivy Severity"] == "high" and r["Clair Severity"] == "medium" and r["Scanners"] == "trivy; grype; clair"
    assert r["Overdue"] == "Yes"
    assert rows[0]["Consensus Severity"] == "critical"


def test_vuln_json(snapshot, opts):
    d = json.loads(generate("vuln-export", "json", snapshot, opts).content)
    assert d["scan"]["id"] == 42 and len(d["findings"]) == sum(len(v[-1]) for v in VULNS)
    f = d["findings"][0]
    assert set(f["perScanner"]) == {"trivy", "grype", "clair"}
    assert f["controls"][:2] == ["RA-5", "SI-2"]


def test_cyclonedx_vex(snapshot, opts):
    rep = generate("vuln-export", "cyclonedx-vex", snapshot, opts)
    d = json.loads(rep.content)
    assert d["bomFormat"] == "CycloneDX" and d["specVersion"] == "1.6"
    refs = {c["bom-ref"] for c in d["components"]}
    assert len(d["vulnerabilities"]) == len({v[0] for v in VULNS})
    for v in d["vulnerabilities"]:
        assert v["analysis"]["state"] == "in_triage"
        assert all(a["ref"] in refs for a in v["affects"])
        assert v["ratings"][0]["severity"] in {"critical", "high", "medium", "low", "info", "unknown"}
