"""Plan of Action & Milestones (POA&M) in eMASS import layouts (xlsx, csv).

Workbook sheets:
  * ``eMASS``          - current eMASS (5.x) POA&M import columns (32, "POA&M Item ID" ...
                         "Resulting Residual Risk after Proposed Mitigations").
  * ``POA&M``          - widely published generic / FedRAMP-style POA&M columns.
  * ``eMASS (legacy)`` - pre-5.x DoD POA&M template columns ("Security Control Number" ...).
  * ``Info``           - system, scan, SLA policy and generation notes.

CSV emits one variant (``options.poamVariant``: ``emass`` (default) | ``generic`` |
``emass-legacy``). Column mappings are documented in docs/REPORTS.md.
"""

from __future__ import annotations

import csv
import io
from copy import copy
from datetime import datetime
from types import SimpleNamespace
from typing import Any

from ._common import (SEVERITIES, TOOL_NAME, View, filename, image_label, mdy, normalize, sev_rank,
                      short_hash)
from .registry import GeneratedReport

XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_CT = "text/csv; charset=utf-8"
MAX_CELL = 32000  # Excel hard limit is 32767 characters per cell

EMASS_COLUMNS = [
    "POA&M Item ID", "Control Vulnerability Description", "Controls / APs", "Security Checks",
    "POA&M Status", "POA&M Scheduled Completion Date", "POA&M Requested Risk Accepted Expiration Date",
    "POA&M Completion Date", "Milestone ID", "Milestone Description", "Milestone Status",
    "Milestone Status Comments", "Milestone Scheduled Completion Date", "Milestone Completion Date",
    "Identification Source", "Identification Source Details", "Office/Org", "Resources Required", "Comments",
    "Raw Severity", "Devices Affected", "Mitigations (in-house and in conjunction with the Navy CSSP)",
    "Predisposing Conditions", "Severity", "Relevance of Threat", "Threat Description", "Likelihood", "Impact",
    "Impact Description", "Residual Risk Level", "Recommendations",
    "Resulting Residual Risk after Proposed Mitigations",
]

GENERIC_COLUMNS = [
    "POAM ID", "Controls", "Weakness Name", "Weakness Description", "Weakness Detector Source",
    "Weakness Source Identifier", "Asset Identifier", "Point of Contact", "Resources Required",
    "Overall Remediation Plan", "Original Detection Date", "Scheduled Completion Date", "Planned Milestones",
    "Milestone Changes", "Status Date", "Vendor Dependency", "Last Vendor Check-in Date",
    "Vendor Dependent Product Name", "Original Risk Rating", "Adjusted Risk Rating", "Risk Adjustment",
    "False Positive", "Operational Requirement", "Deviation Rationale", "Supporting Documents", "Comments",
    "Auto-Approve",
]

LEGACY_COLUMNS = [
    "Control Vulnerability Description", "Security Control Number (NC/NA controls only)", "Office/Org",
    "Security Checks", "Resources Required", "Scheduled Completion Date", "Milestone with Completion Dates",
    "Milestone Changes", "Source Identifying Vulnerability", "Status", "Comments", "Raw Severity",
    "Devices Affected", "Mitigations", "Predisposing Conditions", "Severity", "Relevance of Threat",
    "Threat Description", "Likelihood", "Impact", "Impact Description", "Residual Risk Level",
    "Recommendations", "Resulting Residual Risk after Proposed Mitigations",
]

EMASS_LEVEL = {"critical": "Very High", "high": "High", "medium": "Moderate", "low": "Low",
               "negligible": "Very Low", "unknown": "Low"}
FEDRAMP_RISK = {"critical": "High", "high": "High", "medium": "Moderate", "low": "Low",
                "negligible": "Low", "unknown": "Low"}
CAT = {"critical": "I", "high": "I", "medium": "II", "low": "III", "negligible": "III", "unknown": "III"}


# --------------------------------------------------------------------------- items
def _cut(s: str, n: int = MAX_CELL) -> str:
    s = s or ""
    return s if len(s) <= n else s[: n - 40] + f"\n... [truncated, {len(s) - n + 40} chars omitted]"


def _vuln_items(v: View, rollup: bool) -> list[SimpleNamespace]:
    groups: dict[tuple, list] = {}
    for f in v.open_findings:
        key = (f.vuln_id,) if rollup else (f.image_id, f.vuln_id, f.package)
        groups.setdefault(key, []).append(f)

    items = []
    for key, fs in groups.items():
        f0 = max(fs, key=lambda f: sev_rank(f.severity))
        severity = f0.severity
        first_seen = min(f.first_seen_at for f in fs)
        images = [v.images_by_id.get(f.image_id) for f in fs]
        images = list({id(i): i for i in images if i is not None}.values())
        assets = [image_label(i) for i in images] or [str(f0.image_id)]
        workloads = sorted({w for i in images for w in i.workloads})
        scanners = sorted({s for f in fs for s in f.scanners})
        per_scanner: dict[str, str] = {}
        for f in fs:
            for s, sv in f.per_scanner.items():
                if s not in per_scanner or sev_rank(sv) > sev_rank(per_scanner[s]):
                    per_scanner[s] = sv
        pkgs = sorted({(f.package, f.installed_version, f.fixed_version or "") for f in fs})
        fixable = all(f.fixable for f in fs)
        controls = list(dict.fromkeys(c for f in fs for c in f.controls))
        title = f0.title or f"{f0.vuln_id} in {f0.package}"
        pkg_txt = "; ".join(f"{p} {iv}" + (f" (fixed in {fx})" if fx else " (no fix available)")
                            for p, iv, fx in pkgs)
        if fixable:
            plan = ("Rebuild the affected image(s) with "
                    + "; ".join(f"{p} >= {fx}" for p, _iv, fx in pkgs if fx)
                    + " (or update the upstream base image / chart version), redeploy, and verify closure with a"
                      " rescan.")
            mitigation = "Upgrade " + "; ".join(f"{p} {iv} -> {fx}" for p, iv, fx in pkgs if fx) + "."
        else:
            plan = ("No fixed version is published. Track the vendor advisory, evaluate compensating controls "
                    "(network isolation, least privilege, removal of the package if unused) and request risk "
                    "acceptance if the fix will not arrive before the scheduled completion date.")
            mitigation = "No vendor fix available; compensating controls pending ISSO review."
        agreement = max((f.agreement or 0) for f in fs) if any(f.agreement is not None for f in fs) else None
        cvss = max((f.cvss for f in fs if f.cvss is not None), default=None)
        desc_parts = [title]
        if f0.description and f0.description != title:
            desc_parts.append(f0.description)
        desc_parts.append(f"Affected package(s): {pkg_txt}.")
        if f0.url:
            desc_parts.append(f"Reference: {f0.url}")
        due = v.sla_due(severity, first_seen)
        overdue = v.overdue(severity, first_seen)
        n_enabled = len([s for s in v.scanners if s.enabled]) or 3
        per_txt = ", ".join(f"{s}={per_scanner.get(s, '?')}" for s in scanners)
        agree_txt = f"{len(scanners)}/{n_enabled} scanners ({per_txt})"
        items.append(SimpleNamespace(
            kind="vulnerability",
            poam_id="SP-" + (f"{f0.vuln_id}-{short_hash(*key)}" if not rollup else f"{f0.vuln_id}"),
            source_id=f0.vuln_id, name=f"{f0.vuln_id} ({', '.join(sorted({p for p, _, _ in pkgs}))})",
            title=title, description="\n".join(desc_parts), controls=controls, severity=severity,
            first_seen=first_seen, due=due, overdue=overdue, assets=assets, workloads=workloads,
            scanners=scanners, detector=v.scanner_label(scanners) or "Trivy/Grype/Clair consensus",
            plan=plan, mitigation=mitigation, fixable=fixable,
            vendor_product="" if fixable else ", ".join(sorted({p for p, _, _ in pkgs})),
            agreement=agreement, cvss=cvss, url=f0.url,
            identification_source=f"Vulnerability scan - {TOOL_NAME} ({v.scanner_label(scanners)})",
            comments=" ".join(filter(None, [
                f"[{TOOL_NAME} id SP-{short_hash(*key)}]",
                f"Detected by {agree_txt}.",
                f"CVSS {cvss}." if cvss is not None else "",
                "Single-scanner finding: validate before remediation." if len(scanners) == 1 else "",
                f"SLA {v.sla_days.get(severity)}d from first seen {first_seen:%Y-%m-%d}.",
                "OVERDUE." if overdue else "",
            ])),
            impact=(f"Exploitation of {f0.vuln_id} in {', '.join(sorted({p for p, _, _ in pkgs}))} could "
                    f"compromise the confidentiality, integrity or availability of workloads running "
                    f"{len(assets)} image(s)."),
        ))
    return items


def _posture_items(v: View) -> list[SimpleNamespace]:
    from .stig import rules_for_check  # local import: stig loads YAML lazily

    groups: dict[str, list] = {}
    for r in v.failed_results:
        groups.setdefault(r.check_id, []).append(r)
    items = []
    for check_id, rs in groups.items():
        c = v.check(check_id)
        severity = max((r.severity for r in rs), key=sev_rank)
        first_seen = min((r.first_seen_at for r in rs if r.first_seen_at),
                         default=v.scan.started_at or v.generated_at)
        assets = sorted({r.key + (f" [{r.container}]" if r.container else "") for r in rs})
        stig = rules_for_check(check_id)
        stig_ids = [s["vulnId"] for s in stig]
        details = sorted({r.detail for r in rs if r.detail})
        due = v.sla_due(severity, first_seen)
        overdue = v.overdue(severity, first_seen)
        desc = f"{c.title}: {c.description}".strip().rstrip(":") if c.description else c.title
        desc += f"\nFailing workloads: {len({r.key for r in rs})}."
        if details:
            desc += "\nDetails: " + "; ".join(details[:20]) + (" ..." if len(details) > 20 else "")
        plan = c.remediation or "Update the workload securityContext / manifest to satisfy the check and redeploy."
        sys_note = " (includes Kubernetes system namespaces)" if any(r.system_namespace for r in rs) else ""
        src_stig = ""
        if stig:
            from .stig import benchmark

            b = benchmark(stig[0]["benchmark"])
            src_stig = f"{b['title']} :: Version {b['version']}, {b['releaseInfo']}"
        items.append(SimpleNamespace(
            kind="posture", poam_id=f"SP-CFG-{check_id}", source_id=", ".join([check_id] + stig_ids),
            name=f"Configuration: {c.title}", title=c.title, description=desc, controls=list(c.controls),
            severity=severity, first_seen=first_seen, due=due, overdue=overdue, assets=assets,
            workloads=sorted({r.key for r in rs}), scanners=[], detector=f"{TOOL_NAME} posture check {check_id}",
            plan=plan, mitigation=plan, fixable=True, vendor_product="", agreement=None, cvss=None, url="",
            identification_source=src_stig or f"Configuration review - {TOOL_NAME}",
            comments=" ".join(filter(None, [
                f"[{TOOL_NAME} id SP-CFG-{check_id}]",
                f"Kubernetes posture check '{check_id}' failed on {len(rs)} container(s){sys_note}.",
                f"Mapped STIG rules: {', '.join(stig_ids)}." if stig_ids else "",
                f"SLA {v.sla_days.get(severity)}d from first seen {first_seen:%Y-%m-%d}.",
                "OVERDUE." if overdue else "",
            ])),
            impact=f"Weakened workload isolation ({c.title.lower()}) increases the impact of a container compromise.",
        ))
    return items


def build_items(v: View) -> list[SimpleNamespace]:
    items = _vuln_items(v, bool(v.options.get("rollupByCve"))) + _posture_items(v)
    items.sort(key=lambda i: (-sev_rank(i.severity), i.kind != "vulnerability", i.due or v.generated_at,
                              i.source_id))
    return items


# --------------------------------------------------------------------------- rows
def _assets_text(i: SimpleNamespace) -> str:
    txt = "\n".join(i.assets)
    if i.kind == "vulnerability" and i.workloads:
        txt += "\nUsed by: " + ", ".join(i.workloads)
    return _cut(txt)


def _milestone(i: SimpleNamespace, v: View) -> str:
    if i.kind == "vulnerability" and i.fixable:
        return f"Rebuild image(s) with fixed package, redeploy and confirm by rescan by {mdy(i.due)}."
    if i.kind == "vulnerability":
        return (f"Monitor vendor for fix; document compensating controls or request risk acceptance by "
                f"{mdy(i.due)}.")
    return f"Remediate workload configuration and confirm by rescan by {mdy(i.due)}."


def emass_row(i: SimpleNamespace, v: View, as_date: bool) -> list[Any]:
    d = (lambda x: x.date() if x else None) if as_date else mdy
    poc = ", ".join(filter(None, [v.system.organization, v.system.poc_name, v.system.poc_email]))
    lvl = EMASS_LEVEL[i.severity]
    return [
        "",                                   # POA&M Item ID - assigned by eMASS on import
        _cut(i.description),
        i.controls[0] if i.controls else "",  # eMASS takes one Control / AP per item
        i.source_id,
        "Ongoing",
        d(i.due), "", "",
        1, _milestone(i, v), "Pending", "", d(i.due), "",
        i.identification_source, _cut(i.detector), poc,
        "Existing O&M staff; no additional funding required." if i.fixable else "Vendor fix required.",
        _cut(i.comments + (f" Additional controls: {', '.join(i.controls[1:])}." if len(i.controls) > 1 else "")),
        lvl, _assets_text(i), _cut(i.mitigation), "", lvl,
        "", "", "", "",                       # Relevance of Threat, Threat Description, Likelihood, Impact
        _cut(i.impact), "", _cut(i.plan), "",
    ]


def generic_row(i: SimpleNamespace, v: View, as_date: bool) -> list[Any]:
    d = (lambda x: x.date() if x else None) if as_date else mdy
    poc = ", ".join(filter(None, [v.system.poc_name, v.system.poc_email])) or v.system.organization
    db_dates = [s.db_updated_at for s in v.scanners if s.name in i.scanners and s.db_updated_at]
    return [
        i.poam_id, ", ".join(i.controls), _cut(i.name), _cut(i.description), _cut(i.detector), i.source_id,
        _assets_text(i), poc,
        "Existing O&M staff" if i.fixable else "Vendor fix required",
        _cut(i.plan), d(i.first_seen), d(i.due), _milestone(i, v), "", d(v.generated_at),
        "No" if i.fixable else "Yes",
        d(max(db_dates)) if (db_dates and not i.fixable) else ("" if as_date else ""),
        i.vendor_product, FEDRAMP_RISK[i.severity], "", "No", "No", "No", "",
        f"{TOOL_NAME} scan {v.scan.id} (SAR / vuln-export)", _cut(i.comments), "No",
    ]


def legacy_row(i: SimpleNamespace, v: View, as_date: bool) -> list[Any]:
    d = (lambda x: x.date() if x else None) if as_date else mdy
    org = ", ".join(filter(None, [v.system.organization, v.system.poc_name, v.system.poc_email]))
    return [
        _cut(i.description), ", ".join(i.controls), org, i.source_id,
        "Existing O&M staff" if i.fixable else "Vendor fix required",
        d(i.due), f"1: {_milestone(i, v)}", "", i.identification_source, "Ongoing", _cut(i.comments),
        CAT[i.severity], _assets_text(i), _cut(i.mitigation), "", CAT[i.severity],
        "", "", "", "", _cut(i.impact), "", _cut(i.plan), "",
    ]


VARIANTS = {
    "emass": ("eMASS", EMASS_COLUMNS, emass_row),
    "generic": ("POA&M", GENERIC_COLUMNS, generic_row),
    "emass-legacy": ("eMASS (legacy)", LEGACY_COLUMNS, legacy_row),
}


# --------------------------------------------------------------------------- writers
def _csv(v: View, items: list, variant: str) -> bytes:
    _sheet, cols, rowf = VARIANTS[variant]
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(cols)
    for i in items:
        w.writerow(rowf(i, v, False))
    return buf.getvalue().encode("utf-8-sig")


def _xlsx(v: View, items: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="E7C4FF")
    overdue_fill = PatternFill("solid", fgColor="FDE2E1")
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    date_cols = {"POA&M Scheduled Completion Date", "Milestone Scheduled Completion Date",
                 "Original Detection Date", "Scheduled Completion Date", "Status Date",
                 "Last Vendor Check-in Date"}
    wide = {"Control Vulnerability Description", "Weakness Description", "Devices Affected", "Asset Identifier",
            "Comments", "Overall Remediation Plan", "Recommendations", "Milestone Description",
            "Mitigations (in-house and in conjunction with the Navy CSSP)", "Mitigations", "Impact Description"}

    scratch = wb.create_sheet("_styles")
    styles = {}
    for is_date in (False, True):
        for overdue in (False, True):
            c = scratch.cell(row=1 + len(styles), column=1)
            c.alignment = wrap
            if is_date:
                c.number_format = "mm/dd/yyyy"
            if overdue:
                c.fill = overdue_fill
            styles[(is_date, overdue)] = copy(c._style)
    wb.remove(scratch)

    for variant in ("emass", "generic", "emass-legacy"):
        title, cols, rowf = VARIANTS[variant]
        ws = wb.create_sheet(title)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill, c.alignment = bold, head_fill, Alignment(wrap_text=True, vertical="center")
        # Performance (25k findings x 3 sheets): ws.max_row / ws[r] (via max_column) are O(cells),
        # and assigning style objects per cell hashes them against the workbook registry. Track
        # the row number, fetch cells with ws.cell (a dict lookup) and copy pre-registered
        # StyleArrays instead.
        for r, i in enumerate(items, start=2):
            ws.append(rowf(i, v, True))
            for idx, col in enumerate(cols, start=1):
                cell = ws.cell(row=r, column=idx)
                cell._style = copy(styles[(col in date_cols and bool(cell.value), bool(i.overdue))])
        for idx, col in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = 60 if col in wide else max(14, min(30, len(col) + 2))
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(items) + 1}"

    info = wb.create_sheet("Info")
    sc = {s: 0 for s in SEVERITIES}
    for i in items:
        sc[i.severity] += 1
    rows = [
        ("Report", "Plan of Action & Milestones"),
        ("System / Project Name", v.system.name),
        ("Organization (Office/Org)", v.system.organization),
        ("eMASS System ID", v.system.emass_system_id),
        ("POC", ", ".join(filter(None, [v.system.poc_name, v.system.poc_email, v.system.poc_phone]))),
        ("Scope", v.scope_label),
        ("Scan ID", v.scan.id),
        ("Scan finished", v.scan.finished_at.replace(tzinfo=None) if v.scan.finished_at else ""),
        ("Generated", v.generated_at.replace(tzinfo=None)),
        ("Generated by", TOOL_NAME),
        ("POA&M items", len(items)),
        ("Rolled up per CVE", "Yes" if v.options.get("rollupByCve") else "No"),
        ("Overdue items", sum(1 for i in items if i.overdue)),
        *[(f"Items - {s}", n) for s, n in sc.items() if n],
        *[(f"SLA days - {s}", d) for s, d in v.sla_days.items()],
        ("", ""),
        ("Notes", "Machine-generated starting point for ISSO review; not a substitute for an assessor. "
                  "Sheet 'eMASS' follows the current eMASS POA&M import columns: paste its rows into the "
                  "POA&M import template downloaded from your eMASS system (row 1 here = template header row). "
                  "POA&M Item ID is left blank so eMASS assigns IDs; our stable id is in Comments. "
                  "Risk analysis columns (Relevance of Threat, Likelihood, Impact, Residual Risk) are left "
                  "for the ISSO. Overdue rows are shaded red."),
    ]
    for k, val in rows:
        info.append([k, val])
        info.cell(row=info.max_row, column=1).font = bold
        if isinstance(val, datetime):
            info.cell(row=info.max_row, column=2).number_format = "yyyy-mm-dd hh:mm"
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 100
    info.cell(row=info.max_row, column=2).alignment = wrap

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate(fmt: str, snapshot: Any, options: dict[str, Any]) -> GeneratedReport:
    v = normalize(snapshot, options)
    items = build_items(v)
    if fmt == "xlsx":
        return GeneratedReport(_xlsx(v, items), filename(v, "poam", "xlsx"), XLSX_CT)
    variant = str(options.get("poamVariant") or "emass").lower()
    if variant not in VARIANTS:
        from .registry import UnsupportedReport

        raise UnsupportedReport(f"unknown poamVariant {variant!r}")
    suffix = "" if variant == "emass" else f"-{variant}"
    return GeneratedReport(_csv(v, items, variant), filename(v, f"poam{suffix}", "csv"), CSV_CT)
